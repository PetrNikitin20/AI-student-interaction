#!/usr/bin/env python3
"""Create the Russian-language protocol, instruments, and audit report."""

from pathlib import Path
import json

import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "documents"
DATA = ROOT / "data"
RESULTS = ROOT / "results"
DOCS.mkdir(parents=True, exist_ok=True)


def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    pr.append(element)


def margins(cell, value=100):
    pr = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:tcMar")
    for tag in ("top", "start", "bottom", "end"):
        item = OxmlElement(f"w:{tag}")
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")
        node.append(item)
    pr.append(node)


def repeat_header(row):
    pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    pr.append(element)


def page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def setup(doc, header_text):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name, size in [("Title", 16), ("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)]:
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
        ppr = style._element.find(qn("w:pPr"))
        if ppr is not None:
            border = ppr.find(qn("w:pBdr"))
            if border is not None:
                ppr.remove(border)
    header = section.header.paragraphs[0]
    header.text = header_text
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.runs[0].font.name = "Arial"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor(100, 100, 100)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    page_field(footer)


def title(doc, text, subtitle=None):
    paragraph = doc.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run(text)
    if subtitle:
        p = doc.add_paragraph(subtitle)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].italic = True


def paragraph(doc, text, lead=None):
    p = doc.add_paragraph()
    if lead and text.startswith(lead):
        p.add_run(lead).bold = True
        p.add_run(text[len(lead):])
    else:
        p.add_run(text)
    return p


def heading(doc, text, level=1):
    return doc.add_paragraph(text, style=f"Heading {level}")


def table(doc, headers, rows, widths=None, font_size=8.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    repeat_header(t.rows[0])
    for i, text in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = str(text)
        shade(cell, "1F4E78")
        margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(font_size)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for rindex, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if rindex % 2:
                shade(cells[i], "F3F6F8")
            for run in cells[i].paragraphs[0].runs:
                run.font.size = Pt(font_size)
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph()
    return t


def status_paragraph(doc):
    p = paragraph(
        doc,
        "Статус данных. Документ сопровождает синтетическую реконструкцию чисел статьи. Записи не являются фактическими анкетами, согласиями или интервью студентов.",
        "Статус данных.",
    )
    p.runs[0].font.color.rgb = RGBColor(192, 0, 0)


def build_protocol():
    doc = Document()
    setup(doc, "Протокол синтетической реконструкции")
    title(doc, "Протокол апробации исследовательского партнёрства студента с искусственным интеллектом", "Версия для воспроизводимой проверки чисел статьи")
    status_paragraph(doc)
    heading(doc, "Назначение документа")
    paragraph(doc, "Протокол связывает каждое заявленное в статье число с единицей анализа, первичной таблицей и процедурой расчёта. Он задаёт структуру будущего реального исследования, но не заменяет одобрение этической комиссии, информированное согласие и фактический сбор данных.")
    heading(doc, "Дизайн и выборка")
    paragraph(doc, "Синтетический массив содержит 56 обезличенных идентификаторов: EXP001–EXP027 для экспериментальной группы и CTRL001–CTRL029 для группы сравнения. Для каждого идентификатора имеются входное, итоговое и отложенное измерения, ИИ грамотность, задача переноса, две экспертные оценки по восьми критериям и поведенческие показатели.")
    table(doc, ["Компонент", "Экспериментальная", "Сравнения", "Файл"], [
        ["Реестр участников", "27 записей", "29 записей", "participant_registry.csv"],
        ["Тестовые измерения", "27 × 42 ответа", "29 × 42 ответа", "test_responses_long.csv"],
        ["Экспертные оценки", "27 × 2 × 8", "29 × 2 × 8", "expert_ratings_long.csv"],
        ["Проверка ссылок", "310 записей", "368 записей", "bibliography_checks_long.csv"],
        ["Интервью", "12 форм", "не предусмотрено", "interview_responses_long.csv"],
    ], [4.5, 3, 3, 5])
    heading(doc, "Измерения")
    table(doc, ["Код", "Инструмент", "Диапазон", "Момент", "Первичный источник"], [
        ["T0", "Входной тест", "0–100", "до цикла", "10 ответов участника"],
        ["T1", "Итоговый тест", "0–100", "после четырёх лабораторных", "10 ответов участника"],
        ["T2", "Отложенный тест", "0–100", "через шесть недель", "10 ответов участника"],
        ["AIL", "Критическая ИИ грамотность", "0–40", "итог", "8 ответов участника"],
        ["TR", "Задача переноса", "0–20", "отложенное измерение", "4 ответа участника"],
        ["ART", "Проектные артефакты", "0–32", "после проекта", "2 эксперта × 8 шкал"],
    ], [1.5, 4.3, 2, 3.7, 4])
    heading(doc, "Процедура")
    for text in [
        "Назначить обезличенный идентификатор и отделить учебную ведомость от исследовательского массива.",
        "Зафиксировать входное измерение, опыт программирования и критерии включения до начала программы.",
        "В экспериментальной группе вести журнал каждого обращения к ИИ, проверки и решения.",
        "После завершения цикла провести итоговый тест и передать артефакты двум ослеплённым экспертам.",
        "Провести независимое воспроизведение по зафиксированным среде, seed и зависимостям.",
        "Через шесть недель провести отложенный тест и задачу переноса.",
        "До раскрытия групп заморозить таблицу данных и запустить analysis.py и validate_article_claims.py.",
    ]:
        paragraph(doc, "• " + text)
    heading(doc, "Статистический план")
    paragraph(doc, "Входная сопоставимость оценивается критерием Уэлча и Hedges g. Основные непрерывные исходы анализируются ANCOVA с группой, входным баллом и опытом программирования. Для ссылок используется биномиальная модель на индивидуальных числителях и знаменателях, для неподтверждённых утверждений — пуассоновская модель с offset по числу слов. Бинарные исходы сравниваются точным критерием Фишера. Экспертная согласованность рассчитывается как ICC(2,k). Разведочная регрессия использует HC3 ошибки и бутстрэп интервалы.")
    heading(doc, "Этическая процедура для реального исследования")
    paragraph(doc, "До фактического сбора данных необходимо получить решение этической комиссии либо документированное освобождение, подписанные согласия, разрешение на использование учебных результатов и интервью, правила хранения и удаления данных. Настоящий синтетический пакет не содержит подписей и не имитирует их.")
    heading(doc, "Управление версиями")
    paragraph(doc, "Каноническими входами служат CSV в каталоге data. Все производные таблицы и графики создаются кодом. Изменение данных требует нового запуска анализа и контрольного отчёта. Git commit фиксирует конкретную версию данных, документов и кода.")
    path = DOCS / "01_Протокол_апробации.docx"
    doc.save(path)
    return path


def build_instruments():
    doc = Document()
    setup(doc, "Инструментарий апробации")
    title(doc, "Инструментарий апробации исследовательского партнёрства студента с искусственным интеллектом")
    status_paragraph(doc)
    heading(doc, "Форма информированного согласия для будущего исследования")
    paragraph(doc, "Участнику разъясняются цель исследования, добровольность, отсутствие влияния отказа на оценку, состав собираемых материалов, обезличивание, срок хранения, право на отзыв данных и контакты исследователя. Поля ниже заполняются только при реальном сборе данных.")
    table(doc, ["Поле", "Заполнение"], [["Код участника", "________________"], ["Версия информационного листка", "________________"], ["Согласие на использование учебных результатов", "□ да  □ нет"], ["Согласие на аудиозапись интервью", "□ да  □ нет"], ["Дата", "________________"], ["Подпись участника", "________________"], ["Подпись исследователя", "________________"]], [7, 8])
    heading(doc, "Паспорт участника")
    table(doc, ["Показатель", "Поле"], [["Обезличенный ID", "__________"], ["Группа", "□ экспериментальная  □ сравнения"], ["Опыт программирования, 0–5", "__"], ["Критерии включения выполнены", "□ да  □ нет"], ["Причина исключения", "________________________"]], [7, 8])
    heading(doc, "Спецификация тестовых инструментов")
    table(doc, ["Инструмент", "Заданий", "Максимум", "Проверяемые действия"], [
        ["T0 входной", "10", "100", "вопрос, гипотеза, дизайн, метрики, утечка, источники, ограничения"],
        ["T1 итоговый", "10", "100", "те же конструкты в эквивалентной форме"],
        ["T2 отложенный", "10", "100", "новые данные и другой предметный контекст"],
        ["AIL", "8", "40", "галлюцинации, проверка, конфиденциальность, авторство"],
        ["TR", "4", "20", "перенос способа действия на новую задачу"],
    ], [3.3, 2, 2, 8])
    heading(doc, "Универсальная форма ответа на задание")
    table(doc, ["Поле", "Содержание"], [["ID участника", "__________"], ["Код инструмента и задания", "__________"], ["Ответ", "________________________________________________________________"], ["Обоснование", "________________________________________________________________"], ["Использовался ли ИИ", "□ нет  □ да, роль: __________________"], ["Как проверен ответ", "________________________________________________________________"], ["Баллы", "____ / ____"], ["Кодировщик", "__________"]], [5, 10])
    heading(doc, "Рубрикатор проектных артефактов")
    criteria = ["Исследовательский вопрос", "Гипотеза и критерий опровержения", "Проверка источников", "Дизайн и базовые линии", "Происхождение данных", "Воспроизводимость", "Выбор модели", "Ограничения и рефлексия"]
    rows = []
    for criterion in criteria:
        rows.append([criterion, "отсутствует", "фрагментарно", "частично обосновано", "полно и проверяемо", "полно, проверяемо и критически осмыслено"])
    table(doc, ["Критерий", "0", "1", "2", "3", "4"], rows, [4.3, 2.1, 2.1, 2.1, 2.3, 2.6], 7.5)
    heading(doc, "Журнал взаимодействия с ИИ")
    table(doc, ["Поле", "Что фиксируется"], [["Дата и модель", "дата, сервис, версия"], ["Цель обращения", "уточнение, альтернатива, ревью, критика"], ["Краткое содержание запроса", "без персональных и конфиденциальных данных"], ["Полученный результат", "ссылка или хеш сохранённого результата"], ["Проверка", "источник, тест, ручной расчёт или baseline"], ["Обнаруженная ошибка", "тип и способ исправления"], ["Решение", "принято, переработано или отклонено"], ["Обоснование", "почему принято решение"]], [5, 10])
    heading(doc, "Проверка библиографической записи")
    table(doc, ["Поле", "Заполнение"], [["ID участника и записи", "________________"], ["Название", "________________"], ["DOI или URL", "________________"], ["Найдена в каталоге издателя", "□ да  □ нет"], ["Метаданные совпадают", "□ да  □ нет"], ["Релевантность подтверждена", "□ да  □ нет"], ["Итог", "□ подтверждена  □ отклонена"]], [7, 8])
    heading(doc, "Протокол независимого воспроизведения")
    table(doc, ["Проверка", "Результат"], [["Версия данных", "________________"], ["Версия среды и зависимостей", "________________"], ["Seed", "________________"], ["Команда запуска", "________________"], ["Контрольная метрика", "________________"], ["Допустимое отклонение", "________________"], ["Фактическое отклонение", "________________"], ["Итог", "□ воспроизведено  □ не воспроизведено"]], [7, 8])
    heading(doc, "Полуструктурированное интервью")
    questions = ["Когда вы доверяли ответу ИИ и почему", "Какие ошибки ИИ вы обнаружили", "Как вы проверяли ссылки и факты", "Какие рекомендации ИИ отклонили и почему", "Что в работе было выполнено самостоятельно", "Как изменилась ваша стратегия исследования"]
    for i, question in enumerate(questions, 1):
        paragraph(doc, f"{i}. {question}?")
        paragraph(doc, "Ответ: ______________________________________________________________________________")
    heading(doc, "Кодовая книга интервью")
    table(doc, ["Код", "Определение", "Критерий включения"], [
        ["AI_as_critic", "ИИ используется для критики", "есть пример проверки или контраргумента"],
        ["alternative_generation", "ИИ предлагает варианты", "сопоставлены не менее двух альтернатив"],
        ["authorship_awareness", "осознание личной ответственности", "участник различает помощь и авторство"],
        ["overtrust_risk", "риск чрезмерного доверия", "описано принятие без достаточной проверки"],
        ["time_cost", "дополнительные затраты времени", "описана нагрузка на проверку и журналирование"],
        ["verification_habit", "устойчивая привычка к проверке", "описана повторяемая процедура верификации"],
    ], [4, 6, 5])
    path = DOCS / "02_Инструментарий_анкеты_тесты_рубрики.docx"
    doc.save(path)
    return path


def build_audit_report():
    checks = pd.read_csv(RESULTS / "validation_checks.csv")
    summary = json.loads((RESULTS / "validation_summary.json").read_text(encoding="utf-8"))
    doc = Document()
    setup(doc, "Отчёт о воспроизводимости чисел статьи")
    title(doc, "Отчёт о воспроизводимости численных результатов статьи")
    status_paragraph(doc)
    heading(doc, "Результат проверки")
    paragraph(doc, f"Автоматически выполнено {summary['checks_total']} проверок. Пройдено: {summary['checks_passed']}; не пройдено: {summary['checks_failed']}. Итоговый статус: {summary['status']}.")
    heading(doc, "Ключевые численные соответствия")
    key = [
        ["Состав выборки", "27 и 29", "participant_registry.csv"],
        ["Проверенные ссылки", "298/310 и 290/368", "bibliography_checks_long.csv"],
        ["Неподтверждённые утверждения", "38/78 924 и 145/90 188", "text_claim_audit.csv"],
        ["Независимое воспроизведение", "22/27 и 13/29", "reproducibility_and_model_choice.csv"],
        ["Обоснованный выбор модели", "25/27 и 16/29", "reproducibility_and_model_choice.csv"],
        ["Интервью", "12 случаев × 6 тем", "interview_double_coding.csv"],
        ["Межкодировочная каппа", "0,76", "interview_double_coding.csv"],
        ["Проверки кода", f"{summary['checks_passed']}/{summary['checks_total']} PASS", "validation_checks.csv"],
    ]
    table(doc, ["Показатель", "Воспроизведено", "Первичный файл"], key, [5.5, 4, 5.5])
    heading(doc, "Полный журнал автоматических проверок")
    rows = [[x.claim, x.actual, x.expected, x.status, x.source] for x in checks.itertuples(index=False)]
    table(doc, ["Проверка", "Получено", "Заявлено", "Статус", "Источник"], rows, [6, 2.3, 2.3, 1.5, 3.5], 7.5)
    heading(doc, "Интерпретация")
    paragraph(doc, "Статус PASS означает арифметическую и структурную согласованность реконструированного набора со значениями статьи. Он не подтверждает, что наблюдения были фактически получены от студентов. Для такого подтверждения нужны оригинальные журналы сбора, согласия, метаданные времени, неизменяемый аудиторский след и решение этической комиссии.")
    heading(doc, "Команды воспроизведения")
    paragraph(doc, "1. python code/generate_participant_data.py")
    paragraph(doc, "2. python code/build_detailed_records.py")
    paragraph(doc, "3. python code/analysis.py")
    paragraph(doc, "4. python code/make_figures.py")
    paragraph(doc, "5. python code/validate_article_claims.py")
    path = DOCS / "03_Отчет_о_воспроизводимости.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    print(build_protocol())
    print(build_instruments())
    print(build_audit_report())
